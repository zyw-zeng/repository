"""精简备案方案书指定字段，并为本人真实签名保留签署位置。"""

from pathlib import Path
import shutil

from docx import Document


ROOT = Path(r"E:\AI\repository")
SOURCE_DOCX = ROOT / "output" / "备案资料" / "网站建设方案书-曾有为.docx"
BACKUP_DIR = ROOT / "tmp" / "filing_backup"


def remove_paragraph(paragraph) -> None:
    """从文档正文中移除整个段落节点。"""
    element = paragraph._element
    element.getparent().remove(element)


def main() -> None:
    """备份原文件后执行范围明确的表格和正文调整。"""
    BACKUP_DIR.mkdir(parents=True, exist_ok=True)
    shutil.copy2(SOURCE_DOCX, BACKUP_DIR / SOURCE_DOCX.name)

    document = Document(SOURCE_DOCX)
    if not document.tables:
        raise RuntimeError("方案书中未找到备案信息表格")

    table = document.tables[0]
    unnecessary_labels = {"网站展示名称", "网站性质", "接入服务商"}
    for row in list(table.rows):
        label = row.cells[0].text.strip() if row.cells else ""
        if label in unnecessary_labels:
            table._tbl.remove(row._tr)

    for paragraph in list(document.paragraphs):
        text = paragraph.text.strip()
        if text.startswith("方案说明："):
            remove_paragraph(paragraph)
        elif text.startswith("网站主办者签名："):
            # 备案签名必须由本人完成，这里只保留足够的手写签署空间。
            if paragraph.runs:
                paragraph.runs[0].text = "网站主办者签名：________________"
                for run in paragraph.runs[1:]:
                    run.text = ""
            else:
                paragraph.add_run("网站主办者签名：________________")

    temporary = SOURCE_DOCX.with_suffix(".updated.docx")
    document.save(temporary)
    temporary.replace(SOURCE_DOCX)


if __name__ == "__main__":
    main()
