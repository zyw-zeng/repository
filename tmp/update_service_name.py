"""修正备案方案书中的新增服务名称。"""

from pathlib import Path
import shutil

from docx import Document


ROOT = Path(r"E:\AI\repository")
TARGET_DOCX = ROOT / "output" / "备案资料" / "网站建设方案书-曾有为.docx"
BACKUP_DIR = ROOT / "tmp" / "filing_backup_service_name"
SERVICE_NAME = "zyw的ai小助手"


def main() -> None:
    """备份当前文件，并只修改基本信息表中的服务名称。"""
    BACKUP_DIR.mkdir(parents=True, exist_ok=True)
    shutil.copy2(TARGET_DOCX, BACKUP_DIR / TARGET_DOCX.name)
    document = Document(TARGET_DOCX)
    if not document.tables:
        raise RuntimeError("方案书中未找到备案基本信息表")

    for row in document.tables[0].rows:
        if row.cells[0].text.strip() != "新增服务名称":
            continue
        target = row.cells[1]
        paragraph = target.paragraphs[0]
        if paragraph.runs:
            paragraph.runs[0].text = SERVICE_NAME
            for run in paragraph.runs[1:]:
                run.text = ""
        else:
            paragraph.add_run(SERVICE_NAME)
        break
    else:
        raise RuntimeError("基本信息表中未找到新增服务名称")

    temporary = TARGET_DOCX.with_suffix(".service-name.docx")
    document.save(temporary)
    temporary.replace(TARGET_DOCX)


if __name__ == "__main__":
    main()
