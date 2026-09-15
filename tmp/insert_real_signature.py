"""从本人提供的签名照片提取真实笔迹并嵌入备案方案书。"""

from pathlib import Path
import shutil

from docx import Document
from docx.shared import Cm
from PIL import Image, ImageOps


ROOT = Path(r"E:\AI\repository")
SOURCE_IMAGE = Path(
    r"C:\Users\ADMINI~1\AppData\Local\Temp\codex-clipboard-01768735-2708-435f-bc1d-02574f408a30.jpg"
)
TARGET_IMAGE = ROOT / "output" / "备案资料" / "曾有为-真实签名.png"
TARGET_DOCX = ROOT / "output" / "备案资料" / "网站建设方案书-曾有为.docx"
BACKUP_DIR = ROOT / "tmp" / "filing_backup_signature"


def extract_signature() -> None:
    """仅移除纸张背景并裁边，不生成或改变本人笔迹结构。"""
    source = Image.open(SOURCE_IMAGE).convert("RGB")
    gray = ImageOps.grayscale(source)
    dark_mask = gray.point(lambda value: 255 if value < 95 else 0)
    bounds = dark_mask.getbbox()
    if bounds is None:
        raise RuntimeError("签名照片中未检测到清晰笔迹")

    left, top, right, bottom = bounds
    margin = 35
    crop_box = (
        max(0, left - margin),
        max(0, top - margin),
        min(source.width, right + margin),
        min(source.height, bottom + margin),
    )
    cropped_gray = gray.crop(crop_box)

    # 用原始灰度计算透明度，只留下真实深色笔迹，不对字形进行生成式重绘。
    alpha = cropped_gray.point(
        lambda value: 0 if value >= 145 else min(255, max(0, int((145 - value) * 2.4)))
    )
    signature = Image.new("RGBA", cropped_gray.size, (15, 15, 15, 0))
    signature.putalpha(alpha)
    signature.save(TARGET_IMAGE)


def insert_into_docx() -> None:
    """将签名放在原签署位置，并保留日期与其余正文。"""
    BACKUP_DIR.mkdir(parents=True, exist_ok=True)
    shutil.copy2(TARGET_DOCX, BACKUP_DIR / TARGET_DOCX.name)
    document = Document(TARGET_DOCX)
    for paragraph in document.paragraphs:
        if not paragraph.text.strip().startswith("网站主办者签名："):
            continue
        paragraph.clear()
        paragraph.add_run("网站主办者签名：")
        picture_run = paragraph.add_run()
        picture_run.add_picture(str(TARGET_IMAGE), width=Cm(3.6))
        break
    else:
        raise RuntimeError("方案书中未找到签名位置")

    temporary = TARGET_DOCX.with_suffix(".signed.docx")
    document.save(temporary)
    temporary.replace(TARGET_DOCX)


def main() -> None:
    """依次提取真实签名并更新 Word 文件。"""
    extract_signature()
    insert_into_docx()


if __name__ == "__main__":
    main()
