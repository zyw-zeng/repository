"""实现 DOCX 文档加载与段落结构提取。"""

from pathlib import Path

from docx import Document

from app.knowledge.types import LoadedBlock, LoadedDocument


def load_docx(path: Path) -> LoadedDocument:
    """按标题样式组织 DOCX 段落，并把表格行转换为可检索文本。"""
    source = Document(path)
    blocks: list[LoadedBlock] = []
    current_heading: str | None = None
    current_paragraphs: list[str] = []

    def flush_block() -> None:
        """把当前标题下的非空段落合并为一个加载块。"""
        text = "\n".join(current_paragraphs).strip()
        if text:
            blocks.append(LoadedBlock(text=text, heading=current_heading))

    for paragraph in source.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        if paragraph.style and paragraph.style.name.lower().startswith("heading"):
            flush_block()
            current_paragraphs.clear()
            current_heading = text
        else:
            current_paragraphs.append(text)
    flush_block()

    # MVP 将每一行表格转成制表符分隔文本，保留基本行列关系。
    for table_index, table in enumerate(source.tables, start=1):
        rows = ["\t".join(cell.text.strip() for cell in row.cells) for row in table.rows]
        table_text = "\n".join(row for row in rows if row.strip())
        if table_text:
            blocks.append(
                LoadedBlock(
                    text=table_text,
                    heading=f"表格 {table_index}",
                    metadata={"content_type": "table"},
                )
            )

    return LoadedDocument(source_path=path, file_type="docx", blocks=blocks)
