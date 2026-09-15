"""四种文档加载器及异常文件的单元测试。"""

from pathlib import Path

import pymupdf
import pytest
from docx import Document

from app.knowledge.loaders import DocumentLoadError, load_document


def test_load_text_with_chinese_and_english(tmp_path: Path) -> None:
    """TXT 加载器应同时保留中文和英文内容。"""
    path = tmp_path / "notes.txt"
    path.write_text("个人知识库\nPersonal knowledge base", encoding="utf-8")

    loaded = load_document(path)

    assert loaded.file_type == "txt"
    assert "个人知识库" in loaded.blocks[0].text
    assert "Personal knowledge base" in loaded.blocks[0].text


def test_load_markdown_preserves_headings(tmp_path: Path) -> None:
    """Markdown 加载器应按照标题保存章节边界。"""
    path = tmp_path / "guide.md"
    path.write_text("# 安装\n执行安装命令。\n## 配置\n填写环境变量。", encoding="utf-8")

    loaded = load_document(path)

    assert [block.heading for block in loaded.blocks] == ["安装", "配置"]
    assert "环境变量" in loaded.blocks[1].text


def test_load_docx_preserves_heading_and_table(tmp_path: Path) -> None:
    """DOCX 加载器应提取标题、正文和表格文本。"""
    path = tmp_path / "manual.docx"
    source = Document()
    source.add_heading("使用说明", level=1)
    source.add_paragraph("这是正文内容。")
    table = source.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "名称"
    table.cell(0, 1).text = "个人知识库"
    source.save(path)

    loaded = load_document(path)

    assert loaded.blocks[0].heading == "使用说明"
    assert loaded.blocks[1].metadata["content_type"] == "table"
    assert "个人知识库" in loaded.blocks[1].text


def test_load_pdf_preserves_page_number(tmp_path: Path) -> None:
    """PDF 加载器应按从 1 开始的页码保存正文。"""
    path = tmp_path / "sample.pdf"
    pdf = pymupdf.open()
    page = pdf.new_page()
    page.insert_text((72, 72), "Knowledge Base PDF")
    pdf.save(path)
    pdf.close()

    loaded = load_document(path)

    assert loaded.blocks[0].page_number == 1
    assert "Knowledge Base PDF" in loaded.blocks[0].text


def test_empty_text_is_rejected(tmp_path: Path) -> None:
    """没有正文的文件不能进入后续索引流程。"""
    path = tmp_path / "empty.txt"
    path.write_text("  \n", encoding="utf-8")

    with pytest.raises(DocumentLoadError, match="没有可用于索引"):
        load_document(path)


def test_corrupted_pdf_is_rejected(tmp_path: Path) -> None:
    """损坏的 PDF 应转换成统一的文档解析错误。"""
    path = tmp_path / "broken.pdf"
    path.write_bytes(b"not-a-real-pdf")

    with pytest.raises(DocumentLoadError, match="文档解析失败"):
        load_document(path)
